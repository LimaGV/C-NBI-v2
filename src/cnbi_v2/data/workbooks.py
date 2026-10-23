"""Inspect XLSX structure without saving, recalculating or resolving dependencies."""

from collections import Counter
from dataclasses import dataclass
from datetime import date, datetime, time
from pathlib import Path
from typing import Any
from xml.etree import ElementTree
from zipfile import BadZipFile, ZipFile

from openpyxl import load_workbook
from openpyxl.formula import Tokenizer
from openpyxl.formula.tokenizer import TokenizerError
from openpyxl.utils import get_column_letter
from openpyxl.utils.exceptions import InvalidFileException

from cnbi_v2.data.hashing import sha256_file


@dataclass(frozen=True)
class FormulaCell:
    coordinate: str
    expression: str | None
    external_reference: bool | None


@dataclass(frozen=True)
class ColumnReport:
    letter: str
    header: str | None
    empty_cells: int
    types: dict[str, int]


@dataclass(frozen=True)
class SheetReport:
    name: str
    declared_dimension: str | None
    observed_rows: int
    observed_columns: int
    inspected_cells: int
    inspection_complete: bool
    header_row: int
    columns: list[ColumnReport]
    apparent_duplicate_headers: dict[str, list[str]]
    empty_cells: int
    types: dict[str, int]
    formula_count: int
    formula_cells: list[FormulaCell]
    formula_details_truncated: bool
    external_formula_count: int
    unparsed_formula_count: int


@dataclass(frozen=True)
class ExternalRelationship:
    part: str
    target: str
    relationship_type: str
    is_hyperlink: bool


@dataclass(frozen=True)
class WorkbookReport:
    filename: str
    sha256: str
    size_bytes: int
    source_unchanged: bool
    sheet_names: list[str]
    non_worksheet_names: list[str]
    sheets: list[SheetReport]
    external_package_parts: list[str]
    external_relationships: list[ExternalRelationship]
    external_defined_names: list[str]
    external_reference_assessment: str
    inspection_complete: bool
    formula_values_status: str
    limitations: list[str]


def _external_reference(expression: str) -> bool | None:
    """Detect direct external range operands; text literals are not references."""
    try:
        tokens = Tokenizer(expression).items
    except (TokenizerError, ValueError, IndexError):
        return None
    return any(
        token.type == "OPERAND"
        and token.subtype == "RANGE"
        and "[" in token.value
        and "]" in token.value
        and "!" in token.value
        for token in tokens
    )


def _package_dependencies(
    archive: ZipFile,
) -> tuple[list[str], list[ExternalRelationship], list[str]]:
    names = archive.namelist()
    parts = sorted(
        name
        for name in names
        if (name.startswith(("xl/externalLinks/", "xl/queryTables/")) and name.endswith(".xml"))
        or name == "xl/connections.xml"
    )
    relationships = []
    for name in sorted(names):
        if not name.startswith("xl/") or not name.endswith(".rels"):
            continue
        document = ElementTree.fromstring(archive.read(name))
        for relation in document:
            if relation.get("TargetMode", "").lower() == "external":
                kind = relation.get("Type", "")
                relationships.append(
                    ExternalRelationship(
                        part=name,
                        target=relation.get("Target", ""),
                        relationship_type=kind,
                        is_hyperlink=kind.endswith("/hyperlink"),
                    )
                )
    workbook_xml = ElementTree.fromstring(archive.read("xl/workbook.xml"))
    defined_names = []
    for item in workbook_xml.findall(".//{*}definedName"):
        expression = item.text or ""
        expression = expression if expression.startswith("=") else "=" + expression
        # Direct bracketed workbook references remain evidence if tokenization fails.
        detected = _external_reference(expression)
        if detected or (detected is None and "[" in expression and "]" in expression):
            defined_names.append(item.get("name", ""))
    return parts, relationships, sorted(defined_names)


def _kind(value: object, data_type: str) -> str:
    if value is None:
        return "empty"
    if data_type == "f":
        return "formula"
    if data_type == "e":
        return "error"
    if isinstance(value, bool):
        return "boolean"
    if isinstance(value, (datetime, date, time)):
        return "date_or_time"
    if isinstance(value, (int, float)):
        return "number"
    if isinstance(value, str):
        return "text"
    return "other"


def _header(value: object) -> str | None:
    if value is None:
        return None
    if isinstance(value, (datetime, date, time)):
        return value.isoformat()
    if isinstance(value, (str, bool, int, float)):
        return str(value)
    return None


def _inspect_sheet(
    sheet: Any,
    *,
    header_row: int,
    max_cells: int,
    detail_limit: int,
) -> SheetReport:
    # openpyxl has no bundled type information; Any is limited to this API boundary.
    try:
        declared_dimension = str(sheet.calculate_dimension())
    except ValueError:
        declared_dimension = None
    # Ignore inaccurate declared bounds while keeping them in the report.
    sheet.reset_dimensions()
    column_types: dict[int, Counter[str]] = {}
    headers: dict[int, str | None] = {}
    formula_cells: list[FormulaCell] = []
    rows = columns = inspected_cells = formula_count = external_count = unparsed_count = 0
    complete = True
    for row_number, row in enumerate(sheet.iter_rows(), start=1):
        if inspected_cells + len(row) > max_cells:
            complete = False
            break
        rows = row_number
        columns = max(columns, len(row))
        inspected_cells += len(row)
        for column_number, cell in enumerate(row, start=1):
            value = cell.value
            kind = _kind(value, cell.data_type)
            counts = column_types.setdefault(column_number, Counter())
            if kind != "empty":
                counts[kind] += 1
            if row_number == header_row:
                headers[column_number] = _header(value)
            if kind == "formula":
                formula_count += 1
                text = value if isinstance(value, str) else getattr(value, "text", None)
                expression = text if isinstance(text, str) else None
                external = _external_reference(expression) if expression is not None else None
                external_count += external is True
                unparsed_count += external is None
                if len(formula_cells) < detail_limit:
                    coordinate = f"{get_column_letter(column_number)}{row_number}"
                    formula_cells.append(FormulaCell(coordinate, expression, external))
    column_reports = []
    duplicate_candidates: dict[str, list[str]] = {}
    total_types: Counter[str] = Counter()
    for index in range(1, columns + 1):
        counts = column_types.get(index, Counter()).copy()
        counts["empty"] = rows - sum(counts.values())
        header = headers.get(index)
        letter = str(get_column_letter(index))
        column_reports.append(ColumnReport(letter, header, counts["empty"], dict(counts)))
        total_types.update(counts)
        if header is not None and header.strip():
            duplicate_candidates.setdefault(header.strip().casefold(), []).append(letter)
    duplicates = {key: value for key, value in duplicate_candidates.items() if len(value) > 1}
    return SheetReport(
        name=str(sheet.title),
        declared_dimension=declared_dimension,
        observed_rows=rows,
        observed_columns=columns,
        inspected_cells=inspected_cells,
        inspection_complete=complete,
        header_row=header_row,
        columns=column_reports,
        apparent_duplicate_headers=duplicates,
        empty_cells=total_types["empty"],
        types=dict(total_types),
        formula_count=formula_count,
        formula_cells=formula_cells,
        formula_details_truncated=formula_count > len(formula_cells),
        external_formula_count=external_count,
        unparsed_formula_count=unparsed_count,
    )


def inspect_workbook(
    path: str | Path,
    *,
    header_row: int = 1,
    max_cells: int = 1_000_000,
    detail_limit: int = 100,
) -> WorkbookReport:
    """Inspect .xlsx bytes in read-only mode; limits are reporting controls only.

    Header position is explicit, not scientifically inferred. Counts describe the
    inspected rectangle, including header cells, not the entire Excel grid.
    NOT_DETECTED never certifies that a workbook is self-contained.
    """
    if header_row < 1 or max_cells < 1 or detail_limit < 1:
        raise ValueError("header_row, max_cells and detail_limit must be positive integers")
    source = Path(path)
    if source.suffix.lower() != ".xlsx":
        raise ValueError("Only .xlsx inspection is supported; .xls/.xlsm require a future decision")
    before = sha256_file(source)
    size = source.stat().st_size
    try:
        with source.open("rb") as stream:
            with ZipFile(stream) as archive:
                parts, relationships, defined_names = _package_dependencies(archive)
            stream.seek(0)
            workbook = load_workbook(
                stream,
                read_only=True,
                data_only=False,
                keep_links=False,
            )
            try:
                sheet_names = list(workbook.sheetnames)
                worksheet_names = {sheet.title for sheet in workbook.worksheets}
                non_worksheet_names = [name for name in sheet_names if name not in worksheet_names]
                sheets = [
                    _inspect_sheet(
                        sheet,
                        header_row=header_row,
                        max_cells=max_cells,
                        detail_limit=detail_limit,
                    )
                    for sheet in workbook.worksheets
                ]
            finally:
                workbook.close()
    except (BadZipFile, ElementTree.ParseError, KeyError, InvalidFileException) as error:
        raise ValueError(f"Invalid or unsupported .xlsx workbook: {error}") from error
    if sha256_file(source) != before:
        raise ValueError("Source bytes changed during inspection; discard this report")
    complete = all(sheet.inspection_complete for sheet in sheets)
    detected = bool(
        parts
        or defined_names
        or any(not item.is_hyperlink for item in relationships)
        or any(sheet.external_formula_count for sheet in sheets)
    )
    uncertain = not complete or any(sheet.unparsed_formula_count for sheet in sheets)
    assessment = (
        "DETECTED_UNRESOLVED" if detected else ("INDETERMINATE" if uncertain else "NOT_DETECTED")
    )
    return WorkbookReport(
        filename=source.name,
        sha256=before,
        size_bytes=size,
        source_unchanged=True,
        sheet_names=sheet_names,
        non_worksheet_names=non_worksheet_names,
        sheets=sheets,
        external_package_parts=parts,
        external_relationships=relationships,
        external_defined_names=defined_names,
        external_reference_assessment=assessment,
        inspection_complete=complete,
        formula_values_status="NOT_VERIFIED",
        limitations=[
            "No save, recalculation, link resolution, snapshot or transformation is performed.",
            "Formula results and cached values are not evaluated or certified.",
            "All sheet names are listed; cell statistics cover worksheets only, not chart sheets.",
            "NOT_DETECTED is not proof of independence; dynamic references may evade detection.",
            "Header row is user-selected; duplicate labels use strip/casefold only.",
            "Dimensions are approximate; counts cover the inspected rectangle including headers.",
            f"Per-sheet cell inspection budget: {max_cells}; formula detail limit: {detail_limit}.",
            "Incomplete scans or unsupported formulas cannot establish absence of dependencies.",
        ],
    )
