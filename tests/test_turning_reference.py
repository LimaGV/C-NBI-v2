"""Integrity of real raw evidence and published expected outputs; no model fitting."""

import ast
import csv
import json
import re
from decimal import Decimal
from pathlib import Path

from openpyxl import load_workbook

from cnbi_v2.data import inspect_workbook, sha256_file

ROOT = Path(__file__).resolve().parents[1]
CASE = ROOT / "tests/reference_cases/turning_8responses"
META = ROOT / "data/metadata/turning_application"
RAW = ROOT / "data/raw/turning_application/CCD_original.xlsx"
RAW_SHA256 = "fd9c463efdcd62ff3d7eaa7eec18d9e112543675c30b1f36838beeb38deebf4a"
RESPONSES = ["T", "MTTF", "WR", "Ra", "Rt", "Kp", "ROI", "OEE"]
TERMS = [
    "beta0",
    "beta1",
    "beta2",
    "beta3",
    "beta11",
    "beta22",
    "beta33",
    "beta12",
    "beta13",
    "beta23",
]


def test_golden_coefficients_shape_order_precision_and_completeness() -> None:
    with (CASE / "expected_coefficients.csv").open(newline="", encoding="utf-8") as stream:
        reader = csv.DictReader(stream)
        rows = list(reader)
        assert reader.fieldnames == ["coefficient", *RESPONSES]
    assert len(rows) == 10
    assert [row["coefficient"] for row in rows] == TERMS
    values = [row[name] for row in rows for name in RESPONSES]
    assert len(values) == 80
    assert all(re.fullmatch(r"-?\d+\.\d{3}", value) for value in values)
    assert all(Decimal(value).is_finite() for value in values)


def test_published_summary_has_exactly_original_responses() -> None:
    with (CASE / "expected_model_summary.csv").open(newline="", encoding="utf-8") as stream:
        rows = list(csv.DictReader(stream))
    assert [row["response"] for row in rows] == RESPONSES
    for row in rows:
        for name in ["S", "R2", "R2_adj", "R2_pred"]:
            assert re.fullmatch(r"\d+\.\d{2}", row[name])
        assert row["classification"] in {"Saddle", "Convex"}
        assert row["optimization_direction"] in {"Min", "Max"}
        assert row["residual_assessment"] == "Normal"


def test_reference_metadata_identifies_verified_primary_source() -> None:
    source = (CASE / "source.yaml").read_text(encoding="utf-8")
    assert "doi: 10.1016/j.engappai.2025.112510" in source
    assert "zotero_item_key: SMMX47JT" in source
    assert "verification_status: FULLTEXT_VERIFIED" in source
    assert "table_pdf_page_1_based: 13" in source
    assert "pdf_sha256: 5717f7a90195157830e42e038782f43b2976c77813ccc11f2b3c01e3cc2be70a" in source
    assert "CODING_SCHEME_PENDING_VERIFICATION" in source
    assert 'WR: "1e-2"' in source
    assert 'ROI: "1e-2"' in source
    assert 'OEE: "1e-2"' in source
    assert "rescaling_applied: false" in source


def test_raw_hash_and_byte_identity_are_pinned() -> None:
    assert RAW.stat().st_size == 15427
    assert sha256_file(RAW) == RAW_SHA256
    provenance = (META / "provenance.yaml").read_text(encoding="utf-8")
    assert f"source_sha256_before_copy: {RAW_SHA256}" in provenance
    assert f"source_sha256_after_copy: {RAW_SHA256}" in provenance
    assert f"raw_sha256: {RAW_SHA256}" in provenance


def test_versioned_inspection_covers_all_external_response_formulas() -> None:
    report = json.loads((META / "workbook_inspection.json").read_text(encoding="utf-8"))
    assert report["sha256"] == RAW_SHA256
    assert report["sheet_names"] == ["Plan1"]
    assert report["external_reference_assessment"] == "DETECTED_UNRESOLVED"
    assert report["formula_values_status"] == "NOT_VERIFIED"
    sheet = report["sheets"][0]
    assert (sheet["observed_rows"], sheet["observed_columns"]) == (20, 12)
    assert sheet["formula_count"] == sheet["external_formula_count"] == 152
    assert len(sheet["formula_cells"]) == 152
    assert not sheet["formula_details_truncated"]
    assert sheet["empty_cells"] == 0
    assert report["privacy_redactions"]


def test_real_raw_inspection_preserves_bytes_and_mtime() -> None:
    before = RAW.read_bytes()
    mtime = RAW.stat().st_mtime_ns
    report = inspect_workbook(RAW, detail_limit=1000)
    assert report.source_unchanged
    assert RAW.read_bytes() == before
    assert RAW.stat().st_mtime_ns == mtime
    assert sha256_file(RAW) == RAW_SHA256


def test_original_explicit_ids_and_response_formulas_are_preserved() -> None:
    workbook = load_workbook(RAW, read_only=True, data_only=False, keep_links=False)
    try:
        rows = list(workbook["Plan1"].iter_rows())
        assert [cell.value for cell in rows[0]] == [
            "Exp",
            "cs",
            "f",
            "md",
            *RESPONSES,
        ]
        assert [row[0].value for row in rows[1:]] == list(range(1, 20))
        assert all(cell.data_type == "f" for row in rows[1:] for cell in row[4:])
    finally:
        workbook.close()


def test_historical_raw_is_ineligible_and_only_explicit_literal_revision_is_processed() -> None:
    folder = ROOT / "data/snapshots/turning_application"
    assert sorted(path.name for path in folder.iterdir()) == ["README.md"]
    processed = ROOT / "data/processed/turning_application"
    assert sorted(path.name for path in processed.iterdir()) == [
        "CCD_values_20261023.csv",
        "README.md",
        "original_coded_factors.csv",
    ]
    dataset = (META / "dataset.yaml").read_text(encoding="utf-8")
    assert "status: RAW_NOT_SELF_CONTAINED" in dataset
    assert "snapshot_status: SNAPSHOT_PENDING_EXTERNAL_SOURCE" in dataset
    assert "computational_use_ready: false" in dataset


def test_checksum_manifest_matches_every_listed_artifact() -> None:
    for line in (META / "checksums.sha256").read_text(encoding="utf-8").splitlines():
        if not line or line.startswith("#"):
            continue
        digest, relative = line.split("  ", 1)
        assert sha256_file(ROOT / relative) == digest, relative


def test_production_code_cannot_depend_on_golden_expected_outputs() -> None:
    forbidden = ["reference_cases", "expected_coefficients.csv", "expected_model_summary.csv"]
    for source in (ROOT / "src").rglob("*.py"):
        tree = ast.parse(source.read_text(encoding="utf-8"))
        for node in ast.walk(tree):
            if isinstance(node, ast.Constant) and isinstance(node.value, str):
                assert not any(marker in node.value for marker in forbidden), source
            if isinstance(node, ast.ImportFrom):
                assert not (node.module or "").startswith("tests"), source
            if isinstance(node, ast.Import):
                assert not any(alias.name.startswith("tests") for alias in node.names), source
