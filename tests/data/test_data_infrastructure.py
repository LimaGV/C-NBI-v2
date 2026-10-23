"""Temporary synthetic fixtures validate infrastructure, never scientific data."""

import hashlib
import importlib
import json
import subprocess
import sys
from pathlib import Path
from zipfile import ZipFile

import pytest
from openpyxl import Workbook
from openpyxl.workbook.defined_name import DefinedName

from cnbi_v2.data import inspect_workbook, sha256_file

ROOT = Path(__file__).resolve().parents[2]


@pytest.fixture
def workbook_path(tmp_path: Path) -> Path:
    path = tmp_path / "synthetic-inspection-fixture.xlsx"
    workbook = Workbook()
    sheet = workbook.active
    sheet.title = "Fixture"
    sheet.append(["run_id", "Value", " value ", "note"])
    sheet.append(["fixture-a", 3, None, True])
    sheet.append(["fixture-b", 4, "=B2*2", "=[missing.xlsx]Sheet1!A1"])
    workbook.create_sheet("Metadata").append(["fixture", True])
    workbook.save(path)
    workbook.close()
    return path


def rewrite_member(path: Path, name: str, replacement: bytes) -> None:
    """Alter only a temporary test fixture's ZIP structure."""
    with ZipFile(path) as archive:
        entries = [(info, archive.read(info.filename)) for info in archive.infolist()]
    with ZipFile(path, "w") as archive:
        for info, content in entries:
            archive.writestr(info, replacement if info.filename == name else content)


def test_sha256_known_vector_and_determinism(tmp_path: Path) -> None:
    path = tmp_path / "hash-fixture"
    path.write_bytes(b"abc")
    expected = "ba7816bf8f01cfea414140de5dae2223b00361a396177a9cb410ff61f20015ad"
    assert sha256_file(path) == sha256_file(path) == expected


def test_sha256_multiple_blocks(tmp_path: Path) -> None:
    path = tmp_path / "large-hash-fixture"
    content = b"a" * (1024 * 1024 + 31)
    path.write_bytes(content)
    assert sha256_file(path) == hashlib.sha256(content).hexdigest()


@pytest.mark.parametrize("function", [sha256_file, inspect_workbook])
def test_missing_file(function, tmp_path: Path) -> None:
    with pytest.raises(FileNotFoundError):
        function(tmp_path / "missing.xlsx")


def test_public_utility_import() -> None:
    module = importlib.import_module("cnbi_v2.data")
    assert module.inspect_workbook is inspect_workbook
    assert module.sha256_file is sha256_file


def test_workbook_structure_formulas_and_types(workbook_path: Path) -> None:
    report = inspect_workbook(workbook_path)
    assert report.filename == workbook_path.name
    assert report.size_bytes == workbook_path.stat().st_size
    assert [sheet.name for sheet in report.sheets] == ["Fixture", "Metadata"]
    assert report.sheet_names == ["Fixture", "Metadata"]
    assert report.non_worksheet_names == []
    sheet = report.sheets[0]
    assert (sheet.observed_rows, sheet.observed_columns) == (3, 4)
    assert sheet.inspected_cells == 12
    assert sheet.empty_cells == 1
    assert sheet.types == {"text": 6, "number": 2, "boolean": 1, "formula": 2, "empty": 1}
    assert sheet.apparent_duplicate_headers == {"value": ["B", "C"]}
    assert [column.header for column in sheet.columns] == ["run_id", "Value", " value ", "note"]
    assert [(item.coordinate, item.external_reference) for item in sheet.formula_cells] == [
        ("C3", False),
        ("D3", True),
    ]
    assert report.external_reference_assessment == "DETECTED_UNRESOLVED"
    assert report.formula_values_status == "NOT_VERIFIED"
    assert report.inspection_complete


def test_inspection_preserves_bytes_and_mtime(workbook_path: Path) -> None:
    before = workbook_path.read_bytes()
    mtime = workbook_path.stat().st_mtime_ns
    report = inspect_workbook(workbook_path)
    assert report.source_unchanged
    assert workbook_path.read_bytes() == before
    assert sha256_file(workbook_path) == report.sha256
    assert workbook_path.stat().st_mtime_ns == mtime
    # Also verifies that file handles were closed on Windows.
    renamed = workbook_path.with_name("renamed-fixture.xlsx")
    workbook_path.rename(renamed)
    assert renamed.read_bytes() == before


def test_formula_details_limit_does_not_hide_external_detection(workbook_path: Path) -> None:
    report = inspect_workbook(workbook_path, detail_limit=1)
    assert report.sheets[0].formula_count == 2
    assert len(report.sheets[0].formula_cells) == 1
    assert report.sheets[0].formula_details_truncated
    assert report.sheets[0].external_formula_count == 1
    assert report.external_reference_assessment == "DETECTED_UNRESOLVED"


def test_incomplete_scan_is_not_certified(workbook_path: Path) -> None:
    report = inspect_workbook(workbook_path, max_cells=4)
    assert not report.inspection_complete
    assert report.sheets[0].observed_rows == 1
    assert report.external_reference_assessment == "INDETERMINATE"


def test_explicit_header_row(workbook_path: Path) -> None:
    report = inspect_workbook(workbook_path, header_row=2)
    assert [column.header for column in report.sheets[0].columns] == [
        "fixture-a",
        "3",
        None,
        "True",
    ]


@pytest.mark.parametrize("option", ["header_row", "max_cells", "detail_limit"])
def test_nonpositive_inspection_options(workbook_path: Path, option: str) -> None:
    with pytest.raises(ValueError, match="positive"):
        inspect_workbook(workbook_path, **{option: 0})


def test_invalid_xlsx_is_rejected_without_changes(tmp_path: Path) -> None:
    path = tmp_path / "invalid.xlsx"
    path.write_bytes(b"not an XLSX package")
    before = sha256_file(path)
    with pytest.raises(ValueError, match="Invalid"):
        inspect_workbook(path)
    assert sha256_file(path) == before


def test_unsupported_format_is_not_silently_converted(tmp_path: Path) -> None:
    path = tmp_path / "legacy.xls"
    path.write_bytes(b"fixture")
    with pytest.raises(ValueError, match="Only .xlsx"):
        inspect_workbook(path)
    assert path.read_bytes() == b"fixture"


def test_inaccurate_declared_dimensions_do_not_hide_cells(workbook_path: Path) -> None:
    with ZipFile(workbook_path) as archive:
        xml = archive.read("xl/worksheets/sheet1.xml")
    rewrite_member(
        workbook_path,
        "xl/worksheets/sheet1.xml",
        xml.replace(b'ref="A1:D3"', b'ref="A1:A1"'),
    )
    report = inspect_workbook(workbook_path)
    assert report.sheets[0].declared_dimension == "A1:A1"
    assert report.sheets[0].observed_rows == 3
    assert report.sheets[0].formula_count == 2


def test_external_package_relationships_are_not_resolved(tmp_path: Path) -> None:
    path = tmp_path / "external-package-fixture.xlsx"
    workbook = Workbook()
    workbook.active["A1"] = "fixture"
    workbook.save(path)
    workbook.close()
    relation = (
        '<Relationships xmlns="http://schemas.openxmlformats.org/package/2006/relationships">'
        '<Relationship Id="rId1" Type="http://schemas.openxmlformats.org/officeDocument/'
        '2006/relationships/externalLinkPath" TargetMode="External" '
        'Target="file:///definitely-missing-fixture.xlsx"/></Relationships>'
    )
    with ZipFile(path, "a") as archive:
        archive.writestr("xl/externalLinks/externalLink1.xml", "<externalLink/>")
        archive.writestr("xl/externalLinks/_rels/externalLink1.xml.rels", relation)
    report = inspect_workbook(path)
    assert report.external_package_parts == ["xl/externalLinks/externalLink1.xml"]
    assert report.external_relationships[0].target == "file:///definitely-missing-fixture.xlsx"
    assert report.external_reference_assessment == "DETECTED_UNRESOLVED"


def test_literals_structured_references_and_internal_links_are_not_external(tmp_path: Path) -> None:
    path = tmp_path / "internal-fixture.xlsx"
    workbook = Workbook()
    sheet = workbook.active
    sheet["A1"] = '="[missing.xlsx]Sheet1!A1"'
    sheet["A2"] = "=SUM(Table1[Value])+Sheet2!A1"
    sheet["A3"] = "=[missing.xlsx]Sheet1!A1"
    sheet["A3"].data_type = "s"
    sheet["A4"] = "Reference webpage"
    sheet["A4"].hyperlink = "https://example.org/reference"
    workbook.create_sheet("Sheet2")["A1"] = 1
    workbook.save(path)
    workbook.close()
    report = inspect_workbook(path)
    assert report.external_reference_assessment == "NOT_DETECTED"
    assert report.external_relationships[0].is_hyperlink
    assert report.formula_values_status == "NOT_VERIFIED"


def test_external_defined_name_is_detected_without_formula_cells(tmp_path: Path) -> None:
    path = tmp_path / "named-reference-fixture.xlsx"
    workbook = Workbook()
    workbook.active["A1"] = "fixture"
    workbook.defined_names.add(DefinedName("ExternalInput", attr_text="'[missing.xlsx]Data'!$A$1"))
    workbook.save(path)
    workbook.close()
    report = inspect_workbook(path)
    assert report.sheets[0].formula_count == 0
    assert report.external_defined_names == ["ExternalInput"]
    assert report.external_reference_assessment == "DETECTED_UNRESOLVED"


def test_unparsed_formula_makes_dependency_assessment_indeterminate(tmp_path: Path) -> None:
    path = tmp_path / "unparsed-formula-fixture.xlsx"
    workbook = Workbook()
    workbook.active["A1"] = '=SUM("unterminated)'
    workbook.save(path)
    workbook.close()
    report = inspect_workbook(path)
    assert report.sheets[0].formula_count == 1
    assert report.sheets[0].unparsed_formula_count == 1
    assert report.external_reference_assessment == "INDETERMINATE"


def test_empty_workbook_is_reported(tmp_path: Path) -> None:
    path = tmp_path / "empty-fixture.xlsx"
    workbook = Workbook()
    workbook.save(path)
    workbook.close()
    sheet = inspect_workbook(path).sheets[0]
    assert sheet.observed_rows == sheet.observed_columns == 0
    assert sheet.empty_cells == 0


@pytest.mark.parametrize("filename", ["dataset_template.yaml", "variable_dictionary_template.yaml"])
def test_metadata_templates_exist(filename: str) -> None:
    assert (ROOT / "data" / "metadata" / filename).is_file()


def test_cli_json_report_and_source_preservation(workbook_path: Path) -> None:
    before = sha256_file(workbook_path)
    result = subprocess.run(
        [sys.executable, str(ROOT / "scripts/data/verify_workbook.py"), str(workbook_path)],
        cwd=ROOT,
        capture_output=True,
        text=True,
        check=False,
    )
    assert result.returncode == 0, result.stderr
    report = json.loads(result.stdout)
    assert report["sha256"] == before
    assert report["external_reference_assessment"] == "DETECTED_UNRESOLVED"
    assert sha256_file(workbook_path) == before


def test_cli_missing_file_is_actionable(tmp_path: Path) -> None:
    result = subprocess.run(
        [
            sys.executable,
            str(ROOT / "scripts/data/verify_workbook.py"),
            str(tmp_path / "missing.xlsx"),
        ],
        cwd=ROOT,
        capture_output=True,
        text=True,
        check=False,
    )
    assert result.returncode == 2
    assert not result.stdout
    assert "missing.xlsx" in json.loads(result.stderr)["error"]
    assert "Traceback" not in result.stderr


@pytest.mark.parametrize(
    ("path", "ignored"),
    [
        ("data/README.md", False),
        ("data/raw/README.md", False),
        ("data/raw/project/README.md", False),
        ("data/raw/private.xlsx", True),
        ("data/snapshots/values.xlsx", True),
        ("data/interim/converted.csv", True),
        ("data/processed/unapproved.csv", True),
        ("data/external/benchmark.csv", True),
        ("data/metadata/dataset_template.yaml", False),
        ("data/metadata/schema.json", False),
        ("data/schemas/dataset.json", False),
        ("references/local_library/article.pdf", True),
        ("references/local_library/article.PDF", True),
        ("references/local_library/book.epub", True),
        ("references/local_library/scan.djvu", True),
        ("references/local_library/book.docx", True),
        ("references/local_library/README.md", False),
        ("references/local_library/notes/source.md", False),
        ("references/local_library/source.bib", False),
        ("references/local_library/source.yaml", False),
        ("references/local_library/source.json", False),
        ("references/local_library/source.csv", False),
        ("references/local_library/credentials-local.json", True),
    ],
)
def test_git_publication_policy(path: str, ignored: bool) -> None:
    result = subprocess.run(
        ["git", "check-ignore", "--no-index", "--quiet", path],
        cwd=ROOT,
        check=False,
    )
    assert result.returncode in (0, 1)
    assert (result.returncode == 0) is ignored
